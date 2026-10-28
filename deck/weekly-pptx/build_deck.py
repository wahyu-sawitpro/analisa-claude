"""Weekly sales review (19-25 Oct 2026) on the SawitPRO 2026 template.
Usage: python3 build_deck.py <template.pptx> <data dir> <out.pptx>
Numbers are computed from the six CSV exports; customer names and phone numbers never reach a slide."""
import sys, re, copy
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from barrier import barrier

TPL, DATA, OUT = sys.argv[1:4]
GREEN, YELLOW, ORANGE, INK, GREY = "007200", "FCD102", "FCA502", "141414", "DFDFDF"
MUTED, TINT, PALE, RED = "5C665C", "EEF5EE", "F6F8F6", "B23A3A"
HEAD, BODY = "Plus Jakarta Sans ExtraBold", "Plus Jakarta Sans"
W1S, W2S = pd.Timestamp("2026-10-12"), pd.Timestamp("2026-10-19")

# ---------------------------------------------------------------- data
def num(v): return pd.to_numeric(str(v).replace(",", ""), errors="coerce")
def key(v):
    d = re.sub(r"\D", "", str(v)) if pd.notna(v) else ""
    return re.sub(r"^(62|0)", "", d) or None
def rp(v, d=None):
    v = float(v or 0); a = abs(v)
    if a >= 1e9: s = f"{v/1e9:.2f}".rstrip("0").rstrip(".") + "B"
    elif a >= 1e6: s = (f"{v/1e6:.0f}" if a >= 1e8 else f"{v/1e6:.1f}".rstrip("0").rstrip(".")) + "M"
    elif a >= 1e3: s = f"{v/1e3:.0f}k"
    else: s = f"{v:.0f}"
    return "Rp" + s
def pct(a, b): return round(a / b * 100) if b else 0
def chg(a, b): return f"{'▲' if a >= b else '▼'} {abs(round((a-b)/b*100))}%" if b else "new"

def sales(f):
    s = pd.read_csv(f, dtype=str)
    for c in ["gmv", "qty", "revenue", "discount", "shipping_cost"]: s[c] = s[c].map(num)
    s["d"] = pd.to_datetime(s.sale_order_date, dayfirst=True)
    s["wk"] = (s.d >= W2S).map({True: "W2", False: "W1"})
    s["k"] = s.customer_phone.map(key)
    s["kg"] = s.item_name.str.extract(r"(\d+)\s*[kK]g")[0].astype(float) * s.qty
    return s
def orders(s):
    return s.groupby(["wk", "main_order_no"]).agg(d=("d", "first"), pic=("pic_name", "first"), k=("k", "first"), cust=("customer_name", "first"),
        gmv=("gmv", "sum"), rev=("revenue", "sum"), cat=("product_cat", "first"), so=("main_so_status", "first"), inv=("main_inv_status", "first"),
        cs=("customer_status", "first"), city=("city", "first"), items=("item_name", list), qty=("qty", list)).reset_index()
def wk(s, o, w):
    x, ox = s[s.wk == w], o[o.wk == w]
    return dict(gmv=x.gmv.sum(), orders=len(ox), cust=ox.cust.nunique(), kg=x.kg.sum(), aov=x.gmv.sum() / max(len(ox), 1))

def short(item):
    if re.search(r"benih|topaz", str(item), re.I): return "Topaz seeds"
    if re.search(r"uji|dokter sawit", str(item), re.I): return "Leaf & soil test"
    if re.search(r"pemetaan", str(item), re.I): return "Field mapping"
    t = re.sub(r"\s*\(.*?\)", "", str(item)); t = re.sub(r"\s*-\s*(Topaz \d.*|Egypt|Peru)$", "", t)
    t = re.sub(r"\b\d+\s*(kg|Kg|Liter)\b", "", t); t = t.replace("MOP/KCL ", "KCL ").replace("Canada Cap ", "")
    t = re.sub(r"(\d+-\d+-\d+)(-\d+)?(\s+[\d.]+\s*B)?", r"\1", t)
    return re.sub(r"\s+", " ", t).strip(" -")
def unit(item):
    return "packs" if re.search(r"benih|topaz", item, re.I) else "units" if re.search(r"liter|uji|pemetaan|dokter", item, re.I) else "sacks"

SH, PL = sales(f"{DATA}/Smallholder Sales last two weeks.csv"), sales(f"{DATA}/plantation Sales last two weeks.csv")
SHO, PLO = orders(SH), orders(PL)
sh1, sh2, pl1, pl2 = wk(SH, SHO, "W1"), wk(SH, SHO, "W2"), wk(PL, PLO, "W1"), wk(PL, PLO, "W2")
sho2, plo2 = SHO[SHO.wk == "W2"].sort_values("gmv", ascending=False), PLO[PLO.wk == "W2"].sort_values("gmv", ascending=False)

shr = pd.read_csv(f"{DATA}/CRM smallholder last week rejection reason.csv", dtype=str)
shw = pd.read_csv(f"{DATA}/CRM smallholder last week success story.csv", dtype=str)
plr = pd.read_csv(f"{DATA}/CRM plantation last week rejection reason.csv", dtype=str)
plw = pd.read_csv(f"{DATA}/CRM plantation last week success story.csv", dtype=str)
for d, c in [(shr, "Phone No (08xx)"), (shw, "Phone No (08xx)"), (plr, "Nomor Telfon Leads"), (plw, "Nomor Telfon Leads")]: d["k"] = d[c].map(key)

# Smallholder rejection export: TikTok live viewers, customers asking about an order, real leads
shr["live"] = shr["Detail Pembahasan"].fillna("").str.contains("live", case=False)
shr["svc"] = shr["Respon Pengguna"].eq("Menanyakan Status Pesanan")
LIVE, SVC = shr[shr.live], shr[shr.svc & ~shr.live]
LEADS = shr[~shr.live & ~shr.svc].copy(); LEADS["b"] = LEADS.apply(barrier, axis=1)
def session(t):
    t = t.lower()
    for k, v in [("gobah", "Desa Gobah, Tambang"), ("sei putih", "Desa Sei Putih, Tapung"), ("hangtuah", "Desa Hangtuah, Perhentian Raja"), ("tokling 1", "Tokling 1 Kampar"), ("tokling 2", "Tokling 2 Kampar")]:
        if k in t: return v
    return "Other live"
LIVE = LIVE.assign(sess=LIVE["Detail Pembahasan"].map(session))

# Leads with a logged potential sales or a stated need (read from the CRM fields and the discussion notes, as logged)
NEED = {  # engagement id -> (product asked, need as stated, extra context) ; verified against the export
    "SMEN-42378": ("NPK Pak Tani 7-6-35", "10 t every 6 months", "Total quoted incl. shipping to Aceh Singkil"),
    "SMEN-42734": ("NPK Yaramila 16-16-16", "4 sacks", "Stock confirmed ready; pick-up at warehouse"),
    "SMEN-42745": ("NPK Pak Tani 13-8-27", "20,000 kg (400 sacks)", "Purchase plan being followed up"),
    "SMEN-42381": ("NPK Mutiara 16-16-16 + mix", "~10 t per order", "Farm shop in Bengkulu; wants COD"),
    "SMEN-42310": ("KCL Mahkota, RP Egypt, Urea", "30 sacks each (90 sacks)", "Wants COD and delivery to Rokan Hilir"),
    "SMEN-42436": ("KCL + NPK (4 products)", "1,000 kg (20 sacks)", "Waiting for the mobile shop in Gunung Sahilan"),
    "SMEN-42259": ("Topaz seeds", "300 seeds", "1.4 ha own land; seeds must be in own name"),
}
POT = LEADS.assign(pot=LEADS["Potential GMV"].map(num)).set_index("Engagement ID")
need_rows = []
for i, (prod, need, ctx) in NEED.items():
    r = POT.loc[i]; need_rows.append(dict(id=i, pic=r["Smallholders Team"], b=r["b"], prod=prod, pot=r["pot"] if r["pot"] > 0 else None, need=need, ctx=ctx))
ORDER = ["SMEN-42378", "SMEN-42734", "SMEN-42745", "SMEN-42381", "SMEN-42310", "SMEN-42436", "SMEN-42259"]
need_rows.sort(key=lambda x: ORDER.index(x["id"]))
pot_total = sum(x["pot"] or 0 for x in need_rows)
named_t = 20 + 10 + 4.5 + 1 + 0.2  # 42745, 42381, 42310 (90 sacks), 42436, 42734 (4 sacks)

# Success stories reconciled with sales (phone first; plantation also by lead name and region where the phone differs)
def sh_story():
    rows = []; o2 = SHO[SHO.wk == "W2"]
    for _, w in shw.iterrows():
        m = o2[o2.k == w.k]
        logged = short(w["Product yang diminta customer"]) if pd.notna(w["Product yang diminta customer"]) else "KCL, NPK, RP (in notes)"
        if len(m):  # several orders by one customer: the one whose product matches what was logged
            mm = m[m["items"].map(lambda its: any(short(x) == logged for x in its) or len(its) > 1 and "notes" in logged)]
            o = (mm if len(mm) else m).iloc[0]
            if any(o.main_order_no == r["order"] for r in rows) and len(m) > 1: o = m[~m.main_order_no.isin([r["order"] for r in rows])].iloc[0]
            rows.append(dict(id=w["Engagement ID"], prod=logged, order=o.main_order_no, gmv=o.gmv, note="Matched in sales"))
        else: rows.append(dict(id=w["Engagement ID"], prod=logged, order="–", gmv=None, note="Team put on CRM however sales not reflected yet"))
    return rows
SH_STORY = sh_story()
SH_MATCHED = {r["order"] for r in SH_STORY if r["order"] != "–"}
SH_NOCRM = sho2[~sho2.main_order_no.isin(SH_MATCHED)]

PL_STORY = [  # 5 CRM rows; status read from "Status Leads", potential and notes; checked against the sales export
    dict(id="PSEN-2293", what="Topaz seeds", val=5.385e6, order="#B2C26748", gmv=PLO.set_index("main_order_no").gmv.get("#B2C26748"), note="Matched in sales (same lead and region; phone differs)"),
    dict(id="PSEN-2153", what="4 leaf + 4 soil tests", val=6.396e6, order="–", gmv=None, note="Team put on CRM however sales not reflected yet"),
    dict(id="PSEN-2149", what="KebunPRO", val=16.8e6, order="–", gmv=None, note="Prospect, closing est. 31 Dec; not a sale yet"),
    dict(id="PSEN-2150", what="KebunPRO", val=16.0e6, order="–", gmv=None, note="Prospect, closing est. 31 Dec; not a sale yet"),
    dict(id="PSEN-2223", what="Topaz seeds", val=None, order="–", gmv=None, note="Prospect, no value logged"),
]
PL_NOCRM = plo2[plo2.main_order_no != "#B2C26748"]

def cat_delta(s):
    g = s.groupby(["wk", "product_cat"]).gmv.sum().unstack(0).fillna(0)
    g["d"] = g.get("W2", 0) - g.get("W1", 0); return g.sort_values("d")

# ---------------------------------------------------------------- deck helpers
prs = Presentation(TPL)
SW, SHH = prs.slide_width, prs.slide_height
LAYOUT = next(l for l in prs.slide_layouts if l.name == "CUSTOM_33_1_1")

def rgb(h): return RGBColor.from_string(h)
def text(sl, x, y, w, h, paras, size=10, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY, name=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, m, 0)
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.alignment = align
        if spacing: para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in ((r, {}) if isinstance(r, str) else r for r in runs):
            run = para.add_run(); run.text = t; f = run.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.name = o.get("font", HEAD if o.get("bold", bold) and o.get("head") else font)
            f.color.rgb = rgb(o.get("color", color))
    if name: tb.name = name
    return tb
def box(sl, x, y, w, h, fill, rounded=True, line=None, name=None):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded: s.adjustments[0] = min(0.12, 0.06 / max(min(w, h), 0.06) * 1.0) if min(w, h) > 0 else 0.1
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if name: s.name = name
    return s
NOTES = {
    "Executive summary": "Weeks run Monday to Sunday: 12-18 Oct vs 19-25 Oct (last order 24 Oct). Sales = the gmv column of the sales exports. 'Matching CRM success story' = success-story engagements found in the sales export by phone (Smallholder) or by lead name and region (Plantation).",
    "Smallholder · success story": "Each success-story engagement was matched to a sales order by phone number and the product it logged. Orders without any engagement in the success-story or rejection export are listed on the right.",
    "Smallholder · rejection reasons": "TikTok live viewers = entries whose notes say the account joined a TikTok live; none has a phone number. Customers asking about an order = status 'Menanyakan Status Pesanan'. The barrier for real leads is read from the reason field first, then the discussion notes.",
    "Smallholder · rejection reasons · potential": "Potential sales is the CRM's potential field as logged. Needs are taken from the 'Fertilizer needs (kg)' field or the quantity the lead stated in the notes. Nothing is priced or estimated.",
    "Plantation · success story": "PSEN-2293 matches order #B2C26748 by lead name and region (the phone in the CRM differs from the sales export). PSEN-2153 has no matching order in the sales export.",
}
def new_slide(title, kicker):
    sl = prs.slides.add_slide(LAYOUT)
    if kicker in NOTES: sl.notes_slide.notes_text_frame.text = NOTES[kicker]
    for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
    text(sl, 1.12, 0.2, 8.4, 0.46, [[(title, {"bold": True, "head": True})]], size=16, anchor=MSO_ANCHOR.BOTTOM, name="Title")
    text(sl, 0.42, 0.82, 9.1, 0.2, [[(kicker.upper(), {"bold": True})]], size=7.5, color=GREEN, name="Kicker")
    return sl
def kpi(sl, x, y, w, h, value, label, hot=False):
    box(sl, x, y, w, h, GREEN if hot else PALE, line=None if hot else GREY)
    text(sl, x + 0.14, y + 0.1, w - 0.28, 0.36, [[(value, {"bold": True, "head": True})]], size=17, color="FFFFFF" if hot else GREEN)
    text(sl, x + 0.14, y + 0.48, w - 0.28, h - 0.54, label, size=8, color="FFFFFF" if hot else MUTED)
def card(sl, x, y, w, h, title, sub=None, fill="FFFFFF"):
    box(sl, x, y, w, h, fill, line=GREY)
    text(sl, x + 0.14, y + 0.1, w - 0.28, 0.22, [[(title, {"bold": True})]], size=10, color=INK)
    if sub: text(sl, x + 0.14, y + 0.31, w - 0.28, 0.2, sub, size=7.5, color=MUTED)
    return y + (0.54 if sub else 0.38)
def table(sl, x, y, w, rows, head, widths, size=7.5, row_h=0.22, align=None, hl=None):
    n, m = len(rows) + 1, len(head)
    shp = sl.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n)); t = shp.table
    tot = sum(widths)
    for j, wd in enumerate(widths): t.columns[j].width = Emu(int(Inches(w) * wd / tot))
    for i in range(n):
        t.rows[i].height = Inches(row_h)
        for j in range(m):
            c = t.cell(i, j); val = head[j] if i == 0 else rows[i - 1][j]
            c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid(); c.fill.fore_color.rgb = rgb(GREEN if i == 0 else ("FFF6D6" if hl and hl(i - 1, j) else ("FFFFFF" if i % 2 else PALE)))
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (align and align[j] == "r") else PP_ALIGN.LEFT
            r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = BODY
            r.font.bold = i == 0; r.font.color.rgb = rgb("FFFFFF" if i == 0 else INK)
    return shp
def hbars(sl, x, y, w, items, fmt, color=GREEN, row=0.3, labw=1.9, valw=0.75):
    mx = max(v for _, v, *_ in items) or 1
    for i, it in enumerate(items):
        k, v = it[0], it[1]; c = it[2] if len(it) > 2 else color; yy = y + i * row
        text(sl, x, yy, labw - 0.08, row, k, size=8, anchor=MSO_ANCHOR.MIDDLE)
        bw = (w - labw - valw) * v / mx
        box(sl, x + labw, yy + row * 0.28, w - labw - valw, row * 0.44, "EEF0EE", rounded=False)
        if bw > 0: box(sl, x + labw, yy + row * 0.28, max(bw, 0.03), row * 0.44, c, rounded=False)
        text(sl, x + w - valw + 0.06, yy, valw - 0.06, row, [[(fmt(v), {"bold": True})]], size=8, anchor=MSO_ANCHOR.MIDDLE)
def pairbars(sl, x, y, w, h, groups, fmt, labels=("Week of 12 Oct", "Week of 19 Oct")):
    """Grouped columns: groups = [(name, before, after)], each group scaled on its own (different units)."""
    n = len(groups); gw = w / n; bw = min(0.34, gw * 0.3); base = y + h - 0.32
    for i, (name, a, b) in enumerate(groups):
        mx = max(a, b) or 1; cx = x + gw * i + gw / 2
        for j, (v, col) in enumerate([(a, "B9D7B9"), (b, GREEN)]):
            bh = (h - 0.72) * v / mx; bx = cx - bw - 0.03 + j * (bw + 0.06)
            box(sl, bx, base - bh, bw, max(bh, 0.02), col, rounded=False)
            text(sl, bx - 0.25, base - bh - 0.2, bw + 0.5, 0.18, [[(fmt[i](v), {"bold": True})]], size=7.5, align=PP_ALIGN.CENTER)
        text(sl, x + gw * i, base + 0.05, gw, 0.18, name, size=8, align=PP_ALIGN.CENTER, color=MUTED)
        text(sl, x + gw * i, base + 0.2, gw, 0.16, chg(b, a), size=7.5, align=PP_ALIGN.CENTER, color=GREEN if b >= a else RED, bold=True)
    lx = x
    for lab, col in zip(labels, ["B9D7B9", GREEN]):
        box(sl, lx, y, 0.12, 0.12, col, rounded=False); text(sl, lx + 0.17, y - 0.02, 1.3, 0.16, lab, size=7.5, color=MUTED); lx += 1.45
def diverge(sl, x, y, w, items, row=0.3, labw=1.4):
    mx = max(abs(v) for _, v in items) or 1; mid = x + labw + (w - labw) / 2; half = (w - labw) / 2 - 0.55
    box(sl, mid, y, 0.01, row * len(items), "9AA39A", rounded=False)
    for i, (k, v) in enumerate(items):
        yy = y + i * row; bw = half * abs(v) / mx
        text(sl, x, yy, labw - 0.08, row, k, size=8, anchor=MSO_ANCHOR.MIDDLE)
        if v >= 0:
            box(sl, mid, yy + row * 0.25, max(bw, 0.02), row * 0.5, GREEN, rounded=False)
            text(sl, mid + bw + 0.05, yy, 0.6, row, [[("+" + rp(v), {"bold": True})]], size=7.5, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
        else:
            box(sl, mid - bw, yy + row * 0.25, max(bw, 0.02), row * 0.5, RED, rounded=False)
            text(sl, mid - bw - 0.65, yy, 0.6, row, [[("−" + rp(-v), {"bold": True})]], size=7.5, color=RED, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
def notes(sl, x, y, w, h, items, fill="FFF8DC"):
    box(sl, x, y, w, h, fill)
    text(sl, x + 0.16, y + 0.1, w - 0.32, h - 0.16, [[(b, {"bold": True}), (" " + t, {})] for b, t in items], size=8.5, spacing=3)

# ---------------------------------------------------------------- slides
# Cover: reuse the template cover and closing; drop the reference slides in between
cover = prs.slides[0]
for shp in cover.shapes:
    if not shp.has_text_frame: continue
    t = shp.text_frame.text
    def put(lines, size, bold):
        tf = shp.text_frame; p0 = tf.paragraphs[0]
        for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
        for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
        r0 = p0.runs[0]; r0.text = lines[0]; r0.font.size = Pt(size)
        for ln in lines[1:]:
            np_ = copy.deepcopy(p0._p); tf._txBody.append(np_)
            from pptx.text.text import _Paragraph
            pp = _Paragraph(np_, tf); pp.runs[0].text = ln; pp.runs[0].font.size = Pt(size * 0.62); pp.runs[0].font.bold = False
    if t.startswith("SawitPRO NEW THEME"): put(["Weekly Sales Review"], 34, True)
    elif t.startswith("Design presentation"): put(["Smallholder & Plantation teams"], 12, True)
    elif t.startswith("To be used"): put(["Week of 19–25 Oct 2026, compared with 12–18 Oct", "Sales exports, CRM success stories and rejection reasons · data to 24 Oct 2026"], 12, True)
keep = {prs.slides[0].slide_id, prs.slides[27].slide_id}
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst):
    if int(sid.get("id")) not in keep:
        prs.part.drop_rel(sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")); sldIdLst.remove(sid)

# 2 Executive summary
crm_seen = sum(r["gmv"] or 0 for r in SH_STORY) + 6.12e6
tot2, tot1 = sh2["gmv"] + pl2["gmv"], sh1["gmv"] + pl1["gmv"]
top_sh = sho2.iloc[0]; grosir = SH[(SH.wk == "W2") & (SH.product_cat == "Pupuk Grosir")].gmv.sum()
sl = new_slide(f"{rp(tot2)} in sales (+{pct(tot2-tot1, tot1)}%), carried by a few big orders", "Executive summary")
kpi(sl, 0.42, 1.08, 2.18, 0.86, rp(tot2), f"Total sales, 19–25 Oct ({chg(tot2, tot1)} from {rp(tot1)})", hot=True)
kpi(sl, 2.72, 1.08, 2.18, 0.86, rp(sh2["gmv"]), f"Smallholder, {sh2['orders']} orders ({chg(sh2['gmv'], sh1['gmv'])})")
kpi(sl, 5.02, 1.08, 2.18, 0.86, rp(pl2["gmv"]), f"Plantation, {pl2['orders']} orders ({chg(pl2['gmv'], pl1['gmv'])})")
kpi(sl, 7.32, 1.08, 2.18, 0.86, f"{pct(crm_seen, tot2)}%", "of sales has a matching CRM success story")
ins = [
    ("Smallholder rode on two wholesale KCL orders.", f"{rp(grosir)} of {rp(sh2['gmv'])} ({pct(grosir, sh2['gmv'])}%); orders fell {sh1['orders']} → {sh2['orders']}. Without each week's largest order, sales are down 7%."),
    ("77% of Smallholder “rejections” are TikTok live viewers.", f"{len(LIVE)} of {len(shr)} entries are accounts that joined a live, with no phone number. {len(LEADS)} are real leads."),
    ("Real leads named Rp111M and ~36 t of demand.", f"2 leads logged a potential sales value; 5 stated a volume. Payment terms, mobile-shop routes and shipping cost hold them back."),
    ("Plantation: Topaz seeds are 83% of sales; CRM can't explain it.", f"One buyer placed two identical Rp47.5M orders. All {len(plr)} rejection entries carry one templated reason and no potential."),
]
y = 2.1
for b, t in ins:
    box(sl, 0.42, y + 0.06, 0.1, 0.1, YELLOW, rounded=False)
    text(sl, 0.62, y, 5.6, 0.7, [[(b, {"bold": True})], [(t, {"color": MUTED})]], size=8.8, spacing=1)
    y += 0.72
box(sl, 6.45, 2.08, 3.05, 2.88, GREEN)
text(sl, 6.65, 2.22, 2.7, 0.2, [[("BOTTOM LINE", {"bold": True})]], size=7.5, color=YELLOW)
text(sl, 6.65, 2.46, 2.7, 1.3, [[("This week's sales came from a few large orders the CRM never saw. The pipeline logged this week is thin: work the 7 leads with named volume and stop logging live viewers as leads.", {"bold": True})]], size=10.5, color="FFFFFF")
text(sl, 6.65, 3.95, 2.7, 0.9, f"Collect {rp(179.1e6)} still unpaid on 3 Smallholder orders, and check the duplicate Rp47.5M Topaz order in Plantation.", size=8, color="E6F2E6")

# 3 Smallholder sales vs week before
cd = cat_delta(SH)
sl = new_slide(f"Smallholder {chg(sh2['gmv'], sh1['gmv'])} to {rp(sh2['gmv'])}; one order is {pct(top_sh.gmv, sh2['gmv'])}% of it", "Smallholder · sales vs the week before")
y0 = card(sl, 0.42, 1.08, 4.6, 2.75, "Week of 19 Oct vs week of 12 Oct", "Each pair scaled on its own")
pairbars(sl, 0.6, y0 + 0.02, 4.25, 2.1, [("Sales", sh1["gmv"], sh2["gmv"]), ("Orders", sh1["orders"], sh2["orders"]), ("Customers", sh1["cust"], sh2["cust"]), ("Tonnes", sh1["kg"] / 1000, sh2["kg"] / 1000)],
         [rp, lambda v: str(int(v)), lambda v: str(int(v)), lambda v: f"{v:.0f} t"])
y0 = card(sl, 5.17, 1.08, 4.33, 2.75, "What moved sales", "Change by product category")
diverge(sl, 5.31, y0 + 0.02, 4.05, [(k, v) for k, v in cd["d"].items() if abs(v) >= 1e6][::-1][:6], row=0.36, labw=1.15)
ex2, ex1 = sh2["gmv"] - top_sh.gmv, sh1["gmv"] - SHO[SHO.wk == "W1"].gmv.max()
c1, c2 = set(SHO[SHO.wk == "W1"].cust), set(SHO[SHO.wk == "W2"].cust); exp = round((ex2 - ex1) / ex1 * 100)
notes(sl, 0.42, 3.95, 9.08, 1.05, [
    ("Bigger orders, fewer of them:", f"average order {rp(sh1['aov'])} → {rp(sh2['aov'])}; volume {sh1['kg']/1000:.0f} t → {sh2['kg']/1000:.0f} t."),
    ("One order changes the picture:", f"{top_sh.main_order_no} ({rp(top_sh.gmv)}, 622 sacks KCL Mentari, {top_sh.pic}). Without each week's largest order, sales are {rp(ex1)} → {rp(ex2)} ({'+' if exp >= 0 else '−'}{abs(exp)}%)."),
    ("Khusus Petani almost disappeared:", f"{rp(cd.loc['Khusus Petani','W1'])} → {rp(cd.loc['Khusus Petani','W2'])}; last week's Rp110M Bungo group order did not repeat, and {len(c1 - c2)} of last week's {len(c1)} customers did not reorder."),
])

# 4 Smallholder orders this week
sl = new_slide(f"{len(SH_NOCRM)} of {sh2['orders']} Smallholder orders ({rp(SH_NOCRM.gmv.sum())}) have no CRM record", "Smallholder · orders this week")
rows = []
for _, o in sho2.iterrows():
    prods = ", ".join(dict.fromkeys(short(i) for i in o["items"])); q = int(sum(o.qty)); u = unit(o["items"][0])
    rows.append([o.main_order_no, o.pic, prods if len(prods) < 34 else prods[:32] + "…", f"{q} {u if q != 1 else u[:-1]}", rp(o.gmv), "Partly paid" if o.inv != "PAID" else "Paid",
                 "Completed" if o.so == "COMPLETED" else "Open", "Yes" if o.main_order_no in SH_MATCHED else "No"])
table(sl, 0.42, 1.08, 6.55, rows, ["Order", "PIC", "Products", "Qty", "Sales", "Payment", "Status", "In CRM"], [1.05, 1.55, 1.9, 0.75, 0.7, 0.75, 0.7, 0.55],
      size=7, row_h=0.27, align="lllrrlll", hl=lambda i, j: j == 7 and rows[i][7] == "No")
kpi(sl, 7.12, 1.08, 2.38, 0.9, rp(179.1e6), "still to collect on 3 partly paid orders (Rp333M sales)", hot=True)
kpi(sl, 7.12, 2.08, 2.38, 0.9, f"{(sho2.so != 'COMPLETED').sum()}/{len(sho2)}", "orders still open (not completed)")
kpi(sl, 7.12, 3.08, 2.38, 0.9, f"{pct(grosir, sh2['gmv'])}%", "of sales is Pupuk Grosir (2 wholesale KCL orders)")
text(sl, 7.12, 4.1, 2.38, 0.9, "Wholesale, Topaz and the Mandailing Natal NPK order came from PICs who logged nothing in the CRM this week.", size=8, color=MUTED)

# 5 Smallholder success story
sm = sum(r["gmv"] or 0 for r in SH_STORY)
sl = new_slide(f"CRM covers {rp(sm)} of sales; {rp(SH_NOCRM.gmv.sum())} has no record", "Smallholder · success story")
y0 = card(sl, 0.42, 1.08, 4.48, 2.8, f"CRM success stories ({len(SH_STORY)})", "Checked against the sales export by phone and product")
table(sl, 0.56, y0 + 0.02, 4.2, [[r["id"], r["prod"], r["order"], rp(r["gmv"]) if r["gmv"] else "–", r["note"]] for r in SH_STORY],
      ["CRM ID", "Logged", "Order", "Sales", "Note"], [0.9, 1.25, 0.85, 0.55, 1.15], size=7, row_h=0.3, align="lllrl")
y0 = card(sl, 5.02, 1.08, 4.48, 2.8, f"In sales, not in the CRM ({len(SH_NOCRM)})", "Sales recorded; no CRM engagement logged")
table(sl, 5.16, y0 + 0.02, 4.2, [[o.main_order_no, o.pic, short(o["items"][0]), rp(o.gmv)] for _, o in SH_NOCRM.iterrows()],
      ["Order", "PIC", "Product", "Sales"], [0.95, 1.55, 1.2, 0.6], size=7, row_h=0.3, align="lllr")
notes(sl, 0.42, 4.0, 9.08, 1.0, [
    ("The CRM sees 16% of sales:", f"{rp(sm)} of {rp(sh2['gmv'])}. All 5 stories are Rudi Nirwana's mobile-shop (toko keliling) sales, logged on the day of the sale."),
    ("The big orders are invisible:", f"{rp(SH_NOCRM[SH_NOCRM.pic != 'Rudi Nirwana'].gmv.sum())} of it comes from Muhammad Abdul Wahab, Simon Sutrado Simanjuntak and Joko Prestiwo, who logged no engagement, so we can't see how these customers were won or when they will reorder."),
])

# 6 Smallholder rejection: what the export holds
bc = LEADS.b.value_counts()
sl = new_slide(f"{len(LIVE)} of {len(shr)} “rejections” are TikTok live viewers, not leads", "Smallholder · rejection reasons")
y0 = card(sl, 0.42, 1.08, 4.25, 2.95, "What the rejection export holds", f"{len(shr)} entries, 19–23 Oct")
hbars(sl, 0.56, y0 + 0.05, 3.95, [("TikTok live viewers (no phone)", len(LIVE), "9AA39A"), ("Customers asking about an order", len(SVC), ORANGE), ("Real leads", len(LEADS), GREEN)], lambda v: str(v), row=0.36, labw=1.85, valw=0.4)
ss = LIVE.sess.value_counts()
text(sl, 0.56, y0 + 1.25, 3.95, 1.1, [[("Live sessions logged: ", {"bold": True}), ("; ".join(f"{k} {v}" for k, v in ss.items()) + ". ", {}),
      ("Every viewer is logged as “considering”, by Merita, Deby and Alyssa, with an account name and no phone number.", {})]], size=7.8, color=MUTED)
y0 = card(sl, 4.8, 1.08, 4.7, 2.95, f"Why the {len(LEADS)} real leads didn't buy", "Read from the reason field and the discussion notes")
hbars(sl, 4.94, y0 + 0.02, 4.42, [(k, int(v)) for k, v in bc.items()], lambda v: str(v), row=0.255, labw=2.15, valw=0.35)
notes(sl, 0.42, 4.15, 9.08, 0.85, [
    ("Most real leads never said what they need:", f"{bc.get('Just asking / catalogue sent', 0)} of {len(LEADS)} only asked prices or got the catalogue; 8 are Deby's database calls logged “will think about it”."),
    ("KUD ties come only from database calls:", "all 4 leads tied to a KUD or agent were called from the database list; TikTok leads ask about price, shipping and the warehouse."),
])

# 7 Smallholder rejection potential
sl = new_slide(f"Real leads logged {rp(pot_total)} of potential sales and named ~{named_t:.0f} t more", "Smallholder · rejection reasons · potential")
kpi(sl, 0.42, 1.08, 2.18, 0.86, rp(pot_total), "potential sales logged in the CRM (2 leads)", hot=True)
kpi(sl, 2.72, 1.08, 2.18, 0.86, f"~{named_t:.0f} t", "fertilizer need stated by 5 more leads")
kpi(sl, 5.02, 1.08, 2.18, 0.86, "300", "Topaz seeds requested by 1 lead")
kpi(sl, 7.32, 1.08, 2.18, 0.86, f"{len(LEADS) - len(need_rows)}/{len(LEADS)}", "real leads with no potential or quantity logged")
table(sl, 0.42, 2.08, 9.08, [[x["id"], x["pic"], x["prod"], rp(x["pot"]) if x["pot"] else "–", x["need"], x["b"], x["ctx"]] for x in need_rows],
      ["CRM ID", "PIC", "Product asked", "Potential sales", "Need (as stated)", "What holds it back", "Context"], [0.8, 1.15, 1.45, 0.8, 1.2, 1.55, 2.1],
      size=7, row_h=0.31, align="lllrlll")
text(sl, 0.42, 4.65, 9.08, 0.35, "Potential sales are shown only where the PIC logged it; otherwise the need is shown as the lead stated it (sacks, kg, tonnes or seeds). No value is estimated.", size=7.5, color=MUTED)

# 8 What would unlock them
sl = new_slide("Payment terms, routes and shipping cost hold real leads back", "Smallholder · what would unlock the leads")
cards = [
    ("Payment terms (COD / tempo)", f"{bc.get('Payment terms (COD / tempo)',0)} leads", "A farm shop in Bengkulu (~10 t per order) and a 90-sack order to Rokan Hilir both asked for cash on delivery; one asked for tempo.", "Decide a COD / tempo rule with finance"),
    ("Mobile shop not on their route", f"{bc.get('No mobile-shop route or long wait',0)} leads", "Gunung Sahilan (1,000 kg ready to buy), Labuhanbatu Utara and another village asked when the toko keliling comes; one asked about the Topaz queue.", "Add these villages to next month's route"),
    ("Price or shipping cost", f"{bc.get('Price or shipping cost',0)} leads", "One lead buys NPK Pak Tani at Rp580k elsewhere vs our Rp593k (+2%). Others ask shipping to Siak, Kuansing, Kalimantan and Aceh Singkil.", "Quote landed price incl. shipping up front"),
    ("Wants warehouse pick-up", f"{bc.get('Wants warehouse pick-up',0)} leads", "Asked for the warehouse address to collect themselves; one has 4 sacks confirmed in stock and only needs a reply.", "Send address and pick-up slot the same day"),
]
for i, (h, n, d, a) in enumerate(cards):
    x = 0.42 + (i % 2) * 4.6; y = 1.08 + (i // 2) * 1.5
    box(sl, x, y, 4.48, 1.38, "FFFFFF", line=GREY)
    text(sl, x + 0.16, y + 0.1, 3.0, 0.22, [[(h, {"bold": True})]], size=10)
    box(sl, x + 3.4, y + 0.1, 0.94, 0.24, YELLOW); text(sl, x + 3.4, y + 0.1, 0.94, 0.24, [[(n, {"bold": True})]], size=8, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, x + 0.16, y + 0.4, 4.15, 0.62, d, size=8.2, color=MUTED)
    text(sl, x + 0.16, y + 1.06, 4.15, 0.22, [[("→ " + a, {"bold": True})]], size=8.2, color=GREEN)
notes(sl, 0.42, 4.1, 9.08, 0.9, [
    (f"{len(SVC)} more entries were customers chasing an order, not leads:", "a 20-sack KCL shortfall on three September orders (promised for 22 Oct), a refund query, delivery-location requests and a Topaz pick-up."),
    ("Tied to a KUD or agent (4) and family decisions (2)", "are not ours to fix this week; move them to a later call list."),
])

# 9 Plantation sales
pd_ = cat_delta(PL); eko = PL[(PL.wk == "W2") & (PL.pic_name == "Eko Prastyo")].gmv.sum(); topaz = PL[(PL.wk == "W2") & (PL.product_cat == "Benih")].gmv.sum()
sl = new_slide(f"Plantation {chg(pl2['gmv'], pl1['gmv'])} to {rp(pl2['gmv'])}; one Topaz buyer is {rp(eko)}", "Plantation · sales vs the week before")
y0 = card(sl, 0.42, 1.08, 4.6, 2.75, "Week of 19 Oct vs week of 12 Oct", "Each pair scaled on its own")
pairbars(sl, 0.6, y0 + 0.02, 4.25, 2.1, [("Sales", pl1["gmv"], pl2["gmv"]), ("Orders", pl1["orders"], pl2["orders"]), ("Customers", pl1["cust"], pl2["cust"])], [rp, lambda v: str(int(v)), lambda v: str(int(v))])
y0 = card(sl, 5.17, 1.08, 4.33, 2.75, "What moved sales", "Change by product category")
diverge(sl, 5.31, y0 + 0.02, 4.05, [(k, v) for k, v in pd_["d"].items() if abs(v) > 0][::-1], row=0.4, labw=1.15)
notes(sl, 0.42, 3.95, 9.08, 1.05, [
    ("Check a possible duplicate:", "#B2C26771 and #B2C26772 are the same customer, same day, 50 Topaz packs each (Rp47.5M each) by Eko Prastyo."),
    ("Without that buyer, Plantation is", f"{rp(pl2['gmv'] - eko)} (vs {rp(pl1['gmv'])} the week before). Mapping (Pemetaan Kebun) is new: 2 orders, Rp18M."),
    ("Nothing is completed yet:", f"all {pl2['orders']} orders are still NEW, though all are paid."),
])

# 10 Plantation success story
sl = new_slide("1 Plantation success is in sales; 1 closing isn't reflected yet", "Plantation · success story")
y0 = card(sl, 0.42, 1.08, 5.35, 2.8, f"CRM success stories ({len(PL_STORY)})", "Checked against the sales export")
table(sl, 0.56, y0 + 0.02, 5.07, [[r["id"], r["what"], rp(r["val"]) if r["val"] else "–", r["order"], r["note"]] for r in PL_STORY],
      ["CRM ID", "Logged", "Value in CRM", "Order", "Note"], [0.75, 1.15, 0.75, 0.8, 2.0], size=7, row_h=0.32, align="llrll",
      hl=lambda i, j: j == 4 and "not reflected" in PL_STORY[i]["note"])
y0 = card(sl, 5.9, 1.08, 3.6, 2.8, f"In sales, not in the CRM ({len(PL_NOCRM)})", "Sales recorded; no CRM engagement logged")
table(sl, 6.04, y0 + 0.02, 3.32, [[o.main_order_no, o.pic.split(" ")[0], short(o["items"][0]), rp(o.gmv)] for _, o in PL_NOCRM.iterrows()],
      ["Order", "PIC", "Product", "Sales"], [0.9, 0.7, 1.05, 0.6], size=7, row_h=0.24, align="lllr")
notes(sl, 0.42, 4.0, 9.08, 1.0, [
    ("PSEN-2153 (4 leaf + 4 soil tests, Rp6.4M):", "team put on CRM however sales not reflected yet. Enter the order or correct the CRM."),
    ("The CRM sees 4% of Plantation sales:", f"{rp(PL_NOCRM.gmv.sum())} of {rp(pl2['gmv'])} has no engagement, including Eko Prastyo's {rp(eko)} Topaz buyer. KebunPRO prospects worth Rp32.8M are dated to close by 31 Dec."),
])

# 11 Plantation rejection
pc = plr.groupby(["PIC (Plantation team)", "Sumber Leads"]).size()
sl = new_slide(f"All {len(plr)} Plantation “rejections” are one template, no potential", "Plantation · rejection reasons")
y0 = card(sl, 0.42, 1.08, 4.4, 2.85, "Every entry is the same", f"{len(plr)} entries by Naya Huwaidah and Devi Mariana Naibaho")
facts = [("Status", "Not interested", len(plr)), ("Reason", "No clear interest yet", (plr['Alasan tidak tertarik'] == 'Belum ada ketertarikan yang jelas').sum()),
         ("Location", "Pekanbaru – Payung Sekaki", (plr['Lokasi (kabupatan - kecamatan)'] == 'Kota Pekanbaru - Payung Sekaki').sum()),
         ("Notes", "Empty", plr['Keterangan'].isna().sum()), ("Potential sales", "Empty or Rp0", plr['Potential sales (IDR)'].fillna('0').eq('0').sum())]
table(sl, 0.56, y0 + 0.02, 4.12, [[a, b, f"{c} of {len(plr)}"] for a, b, c in facts], ["Field", "Value logged", "Entries"], [1.0, 1.9, 0.9], size=7.5, row_h=0.32, align="llr")
y0 = card(sl, 4.95, 1.08, 4.55, 2.85, "Who logged them, and for what", "Entries by PIC and lead source; product offered")
hbars(sl, 5.09, y0 + 0.02, 4.27, [(f"{p} · {s}", int(v)) for (p, s), v in pc.sort_values(ascending=False).items()], lambda v: str(v), row=0.3, labw=2.6, valw=0.35)
pu = plr['Tujuan Melakukan Engagement'].value_counts()
text(sl, 5.09, y0 + 1.35, 4.27, 0.8, [[("Offered: ", {"bold": True}), (", ".join(f"{k} {v}" for k, v in pu.items()) + ". About 10 entries per PIC per day.", {})]], size=7.8, color=MUTED)
notes(sl, 0.42, 4.05, 9.08, 0.95, [
    ("Potential can't be shown:", "no entry has a potential sales value, an FFB volume or a stated need, so the size of what was lost is unknown."),
    ("Yet Topaz sells:", f"Topaz is {pct(topaz, pl2['gmv'])}% of Plantation sales this week, while 51 Topaz ad leads are logged as not interested. Devi Mariana Naibaho sold a Rp9.5M Topaz order that has no CRM entry."),
])

# 12 Next steps
sl = new_slide("Six actions for the coming week", "Next steps")
acts = [
    ("Work the 7 leads with named demand", "Rp111M logged + ~36 t stated (table on slide 7). Call each this week; send the warehouse slot to the 4-sack pick-up today.", "Merita, Alyssa, Aisyah, Deby"),
    ("Decide COD / tempo with finance", "Two volume buyers (~10 t store in Bengkulu, 90 sacks to Rokan Hilir) wait on cash on delivery. Agree a rule and reply.", "Sales lead + Finance"),
    ("Collect Rp179M still unpaid", "3 Smallholder orders are partly paid, led by the Rp267M wholesale KCL order.", "Wahab, Joko, Rudi + Finance"),
    ("Log leads, not live viewers", "Keep TikTok live viewers in a separate count; log a lead only with a phone number and a stated need (kg or sacks).", "Merita, Deby, Alyssa · Sales ops"),
    ("Put every sale in the CRM", "Rp432M (Smallholder) and Rp134M (Plantation) has no engagement. Enter PSEN-2153's order; check the duplicate Topaz order.", "All PICs · Sales ops"),
    ("Replace Plantation's templated entries", "Record the real reason, location and need for ads and event leads, starting with the 51 Topaz leads.", "Naya, Devi"),
]
for i, (h, d, o) in enumerate(acts):
    x = 0.42 + (i % 3) * 3.05; y = 1.08 + (i // 3) * 1.95
    box(sl, x, y, 2.93, 1.83, "FFFFFF", line=GREY)
    box(sl, x + 0.14, y + 0.14, 0.32, 0.32, GREEN); text(sl, x + 0.14, y + 0.14, 0.32, 0.32, [[(str(i + 1), {"bold": True})]], size=10, color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, x + 0.56, y + 0.12, 2.25, 0.4, [[(h, {"bold": True})]], size=9.5)
    text(sl, x + 0.14, y + 0.6, 2.65, 0.85, d, size=8, color=MUTED)
    box(sl, x + 0.14, y + 1.45, 2.65, 0.26, TINT); text(sl, x + 0.2, y + 1.45, 2.55, 0.26, [[("Owner: ", {"bold": True}), (o, {})]], size=7.5, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)

# move the closing slide (template "Terima Kasih") to the end
lst = prs.slides._sldIdLst; ids = list(lst)
closing = [s for s in ids if int(s.get("id")) in keep and s is not ids[0]][0]
lst.remove(closing); lst.append(closing)
prs.save(OUT)
print("slides", len(prs.slides))
